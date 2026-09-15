"""Genera el documento de trabajo del equipo con los 100 casos de uso de IA.

    data/casos-ia.json  (que hace cada caso y sus dos etiquetas)
    data/domains/*.json (donde aparece cada uno)
        -> Casos_de_uso_IA_Finance_Strategy.xlsx

Esta es la direccion del flujo, y es la misma que la de todo lo demas aqui: el
repositorio es la fuente y el documento es la salida. El Excel no se versiona
porque se puede volver a generar en cualquier momento; lo que se versiona es de
donde sale.

El generador anterior vivia en un directorio temporal y se perdio, que es lo que
pasa con los scripts que no entran en el repositorio.

    python scripts/generar_excel_casos_ia.py [ruta-de-salida]

Dos avisos sobre lo que este script NO reproduce del documento original:

- "Dominio principal" y "Capacidad principal" se recalculan con una regla
  explicita —donde mas veces aparece el caso, y a igualdad el orden del
  catalogo—, porque la del documento original no se conserva en ningun sitio y
  no era deducible: habia casos de un solo dominio etiquetados "Transversal".
  La hoja de resumen se agrupa por esa columna, asi que alguna cifra puede
  moverse respecto a la copia que tenga el equipo. Los totales siguen sumando
  100.

- La columna "Notas del equipo" sale vacia. Es texto libre que no vive en el
  repositorio, asi que una regeneracion la pisa. Antes de regenerar, avisa.
"""

import json
import sys
from collections import Counter, OrderedDict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]

SALIDA_POR_DEFECTO = ROOT / "Casos_de_uso_IA_Finance_Strategy.xlsx"

AZUL_DELOITTE = "012169"
VERDE_DELOITTE = "86BC25"

HOJA_CASOS = "2. Los 100 casos"
HOJA_LISTAS = "Listas"

CABECERAS = [
    ("ID", 9),
    ("Caso de uso de IA", 42),
    ("Qué hace", 74),
    ("Tipo de IA", 20),
    ("Tipo de valor", 22),
    ("Dominio principal", 21),
    ("Capacidad principal", 30),
    ("Subcapacidades", 13),
    ("Dominios", 10),
    ("Dónde aparece", 40),
    ("En qué subcapacidades", 54),
    ("POC en cliente", 13),
    ("Origen", 48),
    ("Notas del equipo", 32),
]

# Texto editorial del documento: lo que explica como usarlo y el registro de la
# limpieza que dejo la lista en 100. No sale de los datos porque no es un dato;
# es lo que el equipo escribio alrededor de ellos. Vive aqui, y no en
# data/casos-ia.json, porque ese archivo lo descarga la aplicacion en cada carga
# y esto no lo lee nadie mas que este script.
COMO_USAR = [
    ("Casos de uso de IA · Finance Strategy", None),
    (None, None),
    (
        "Qué es",
        "La lista cerrada de los 100 casos de uso de IA de F3M. Es el documento de trabajo "
        "del equipo y el mismo contenido que muestra la herramienta.",
    ),
    (
        "Cómo se mantiene",
        "Este documento es una salida, no una fuente. Los títulos y su reparto por "
        "subcapacidad vienen de los nueve Excel de dominio; «Qué hace» y las dos etiquetas "
        "viven en data/casos-ia.json. Se regenera con "
        "'python scripts/generar_excel_casos_ia.py'. Si corriges algo aquí, llévalo al "
        "repositorio: la siguiente regeneración pisa el cambio.",
    ),
    (
        "Qué enseña la herramienta",
        "Los mismos 100 casos, ya clasificados: cada ficha lleva su frase de «qué hace» y "
        "sus dos etiquetas, en el detalle de cada subcapacidad y en el modal de IA del "
        "roadmap.",
    ),
    (None, None),
    ("Las dos etiquetas", None),
    (
        "Tipo de IA",
        "Qué tecnología hay detrás. Sirve para no vender como IA lo que es automatización, "
        "y para saber cuándo hablamos de agentes.",
    ),
    (
        "Tipo de valor",
        "Qué palanca mueve el caso. Es la que ordena la conversación con el cliente: "
        "¿quieres atacar coste, riesgo, decisión o P&L?",
    ),
    (None, None),
    ("Cómo usarlo en una conversación", None),
    ("Filtra por dominio", "Para preparar una sesión de un dominio concreto."),
    ("Filtra por tipo de valor", "Cuando el cliente ya ha dicho qué le duele."),
    (
        "Filtra por POC en cliente",
        "Los que tienen un POC vivo o en curso según el catálogo de FinanceAI. Son los que "
        "se pueden contar como hechos, no como propuesta.",
    ),
    (
        "Mira «Dónde aparece»",
        "Un caso que sale en cinco dominios es plataforma, no un caso de un dominio. Suelen "
        "ser los habilitadores transversales.",
    ),
    (
        "«En qué subcapacidades»",
        "Los identificadores exactos, en el orden del catálogo. Sirven para ir directo al "
        "sitio en la herramienta y en los Excel de origen: el formato es "
        "dominio-capacidad-subcapacidad, así que «fiscal-2-3» es la 2.3 de Fiscal.",
    ),
    (None, None),
    ("Una cautela con dos columnas", None),
    (
        "Dominio y capacidad principal",
        "Se calculan: donde más veces aparece el caso, y a igualdad el orden del catálogo de "
        "dominios. Son un resumen, no una asignación del equipo. Lo que manda es «Dónde "
        "aparece».",
    ),
    (
        "Notas del equipo",
        "Columna libre que no vive en el repositorio. Cada regeneración la deja vacía: si "
        "hay notas que conservar, cópialas antes.",
    ),
]

# El registro de la limpieza que cerro la lista en 100. Es historia y no cambia:
# se conserva para que el documento siga explicando por que hay 100 casos y no
# los 101 de antes.
CAMBIOS_APLICADOS = [
    ("Qué", "Caso", "Motivo"),
    (
        "Retirado",
        "IA-069 — Optimización de red y eficiencia energética",
        "Operación de red eléctrica. En «riesgo de tipos y contraparte» no tiene lectura "
        "financiera: el vínculo era la palabra «optimización».",
    ),
    (
        "Retirado",
        "IA-092 — Operaciones de tienda de nueva generación",
        "Coordinación de la actividad en tienda. No toca ningún proceso financiero.",
    ),
    (
        "Retirado",
        "IA-093 — Gestión de permisos y licencias",
        "Tramitación de permisos y licencias administrativas.",
    ),
    (
        "Retirado",
        "IA-094 — Agentes de IA para el ciclo de vida del servicio",
        "Ciclo de vida de servicios de telecomunicaciones.",
    ),
    (
        "Retirado",
        "IA-095 — Gestión autónoma de operaciones de campo",
        "Gestión del trabajo de campo de técnicos.",
    ),
    (
        "Retirado",
        "IA-097 — Inspección de activos con drones",
        "Inspección de infraestructura con drones. Vínculo financiero indirecto.",
    ),
    (
        "Retirado",
        "IA-109 — Asistente de compra y leasing de vehículos",
        "Venta minorista de automoción al consumidor final.",
    ),
    (
        "Retirado",
        "IA-125 — Resolución autónoma de garantías",
        "Resolución de garantías de automoción. La analogía era solo terminológica.",
    ),
    (
        "Retirado",
        "IA-126 — Inspección visual de partículas",
        "Visión artificial para detectar partículas en producto farmacéutico.",
    ),
    (
        "Retirado",
        "IA-068 — Análisis predictivo con IA generativa",
        "No es un caso de uso: «Análisis predictivo + GenAI» es un valor de la columna «Tipo "
        "de IA» del catálogo maestro, colado en la lista de casos.",
    ),
    (
        "Fusionado",
        "IA-011 — Conciliación asistida por IA → Conciliación a nivel de transacción",
        "«Conciliación con IA» era la versión genérica de la conciliación a nivel de "
        "transacción. No se diferenciaban.",
    ),
    (
        "Añadido",
        "IA-135 — Generación de contenido para el earnings call",
        "Slide FinanceAI «What AI Use Cases exist for Finance?», Investor Relations. Entra "
        "en Relación con Inversores / Planificación y estrategia de RI.",
    ),
    (
        "Añadido",
        "IA-136 — Reporting estandarizado de revisión de desempeño",
        "Slide FinanceAI, Strategic Finance. Entra en FP&A / Informes de gestión del "
        "rendimiento.",
    ),
    (
        "Añadido",
        "IA-137 — Detección de facturas duplicadas",
        "Slide FinanceAI, Source-to-Pay. Entra en Transacciones / Source-to-Pay.",
    ),
    (
        "Añadido",
        "IA-138 — Insights sobre el dato de proveedor",
        "Slide FinanceAI, Source-to-Pay. Entra en Transacciones / Source-to-Pay.",
    ),
    (
        "Añadido",
        "IA-139 — Modelos retadores para contrastar escenarios",
        "Catálogo maestro F3M_AI_Mapping_Consolidado_v1, fila 39, encaje Alto. Entra en "
        "FP&A / Información y apoyo a la toma de decisiones.",
    ),
    (
        "Añadido",
        "IA-140 — Asistente financiero para el cliente",
        "Catálogo maestro, fila 82; dossier p. 85, Financial Services. Entra en "
        "Transacciones / Order-to-Cash.",
    ),
    (
        "Añadido",
        "IA-141 — Respuesta asistida a requerimientos fiscales",
        "Slide FinanceAI, Tax. Entra en Fiscal / Cumplimiento tributario directo.",
    ),
    (
        "Añadido",
        "IA-142 — Borrador de informe de auditoría",
        "Slide FinanceAI, Internal Audit. Entra en Auditoría Interna / Aseguramiento de "
        "riesgos.",
    ),
    (
        "Añadido",
        "IA-143 — Generación de casos de prueba de control",
        "Dossier p. 183, Technology, Media & Telecommunications. Entra en Auditoría Interna "
        "/ Aseguramiento de riesgos.",
    ),
    (
        "Añadido",
        "IA-144 — Determinación de elegibilidad para incentivos",
        "Dossier p. 101, Government & Public Services. Entra en Finanzas de negocio / "
        "Financiación I+D.",
    ),
    (
        "Descartado",
        "Venta técnica asistida por IA",
        "Catálogo maestro, fila 101. Solapa con el asistente de RFP, que ya está en la misma "
        "capacidad.",
    ),
    (
        "Descartado",
        "Revisión automatizada de expedientes de reclamación",
        "Catálogo maestro, fila 84. El caso original es de seguros; su traslación a una "
        "función financiera corporativa tiene confianza baja.",
    ),
    (
        "Descartado",
        "Creación de modelos analíticos por el propio equipo",
        "Dossier p. 147. Describe una capacidad del equipo, no un caso de uso: encaja mejor "
        "como criterio de madurez.",
    ),
    (
        "Descartado",
        "Análisis de escenarios a nivel de proyecto",
        "Slide FinanceAI, Capital Planning. Solapa con «Recomendaciones de acción a nivel de "
        "proyecto» y con «Previsión y escenarios bajo demanda».",
    ),
    (
        "Descartado",
        "Generación de varios modelos de escenario",
        "Catálogo maestro, fila 38. Solapa con «Previsión y escenarios bajo demanda», que ya "
        "está en la herramienta.",
    ),
    (
        "Descartado",
        "Hub de insights ejecutivo con tendencias y KPIs",
        "Slide FinanceAI, Strategic Finance. Solapa con el analista digital y con la visión "
        "ejecutiva de salud financiera.",
    ),
    (
        "Unificado",
        "31 redacciones redundantes",
        "Cada caso pasa a tener un único título. El analista digital tenía cinco nombres "
        "distintos; las estrategias de inversión, tres longitudes.",
    ),
    (
        "Reasignado",
        "11 subcapacidades reforzadas",
        "Se quedaban con un solo caso. Se les ha asignado un segundo caso del catálogo "
        "existente, sin inventar ninguno.",
    ),
]

CRITERIOS_Y_ADVERTENCIA = [
    (None, None),
    ("Criterios de selección de los 100", None),
    (
        "Se retiró",
        "Lo que no describe una capacidad de la función financiera, por muy AI que sea: "
        "drones, tienda, permisos, garantías de automoción, red eléctrica.",
    ),
    ("Se fusionó", "El caso genérico cuando ya existía el específico."),
    ("Se unificó", "Un solo título por caso. Antes había 132 redacciones para 101 casos."),
    (
        "Se añadió",
        "Priorizando, por este orden: que tenga POC vivo en cliente, que el catálogo maestro "
        "le dé encaje alto, y que refuerce un dominio flojo.",
    ),
    (
        "Se descartó",
        "Lo que solapaba con un caso ya presente, y lo que describe una capacidad del equipo "
        "más que un caso de uso. Está en la pestaña 4, con el motivo.",
    ),
    (None, None),
    ("Una advertencia", None),
    (
        None,
        "«Tipo de valor» es una clasificación de criterio, no un dato heredado de ninguna "
        "fuente. Se ha asignado caso a caso y es discutible en los límites: un caso de "
        "conciliación puede leerse como eficiencia o como control. Si el equipo la cambia, "
        "que quede en data/casos-ia.json y se regenere desde ahí.",
    ),
]


def leer_catalogo():
    return json.loads((ROOT / "data" / "casos-ia.json").read_text(encoding="utf-8"))


def leer_apariciones():
    """Donde aparece cada caso: (id de dominio, capacidad, id de subcapacidad).

    Se recorre en el orden del catalogo de dominios para que "Dónde aparece" y
    "En qué subcapacidades" salgan siempre igual entre dos ejecuciones.
    """
    catalogo = json.loads((ROOT / "data" / "domains.json").read_text(encoding="utf-8"))

    etiquetas = OrderedDict((d["id"], d["label"]) for d in catalogo["domains"])
    apariciones = {}

    for dominio in catalogo["domains"]:
        datos = json.loads((ROOT / dominio["dataUrl"]).read_text(encoding="utf-8"))

        for item in datos["subcapacities"]:
            cases = (item.get("ai") or {}).get("cases", "")

            for titulo in [t.strip() for t in cases.split(";") if t.strip()]:
                apariciones.setdefault(titulo, []).append(
                    (dominio["id"], item["capacidad"], item["id"])
                )

    return etiquetas, apariciones


def principal(ocurrencias, orden):
    """Dominio y capacidad donde mas veces aparece el caso.

    A igualdad de apariciones gana el que va antes en el catalogo de dominios,
    para que el resultado no dependa del orden de un diccionario.
    """
    if not ocurrencias:
        return "", ""

    por_dominio = Counter(o[0] for o in ocurrencias)

    dominio = max(por_dominio, key=lambda did: (por_dominio[did], -orden.index(did)))

    por_capacidad = Counter(o[1] for o in ocurrencias if o[0] == dominio)

    capacidad = max(por_capacidad, key=lambda cap: por_capacidad[cap])

    return dominio, capacidad


def estilar_cabecera(hoja, fila=1):
    for celda in hoja[fila]:
        celda.font = Font(bold=True, color="FFFFFF", size=10)
        celda.fill = PatternFill("solid", fgColor=AZUL_DELOITTE)
        celda.alignment = Alignment(vertical="center", wrap_text=True)

    hoja.row_dimensions[fila].height = 32


def escribir_como_usar(libro):
    hoja = libro.active
    hoja.title = "1. Cómo usar este documento"

    for etiqueta, texto in COMO_USAR:
        hoja.append([etiqueta, texto])

    hoja["A1"].font = Font(bold=True, size=15, color=AZUL_DELOITTE)

    for fila in hoja.iter_rows(min_row=2):
        etiqueta, texto = fila[0], fila[1]

        # Los titulos de bloque van solos en su fila: en negrita y en verde.
        if etiqueta.value and not texto.value:
            etiqueta.font = Font(bold=True, size=11, color="3E6F11")
        elif etiqueta.value:
            etiqueta.font = Font(bold=True, size=10)

        texto.alignment = Alignment(wrap_text=True, vertical="top")

    hoja.column_dimensions["A"].width = 30
    hoja.column_dimensions["B"].width = 116


def escribir_casos(libro, catalogo, etiquetas, apariciones):
    hoja = libro.create_sheet(HOJA_CASOS)

    hoja.append([titulo for titulo, _ in CABECERAS])

    orden = list(etiquetas)

    for caso in catalogo["casos"]:
        ocurrencias = apariciones.get(caso["titulo"], [])

        dominios = list(OrderedDict.fromkeys(o[0] for o in ocurrencias))
        dominio, capacidad = principal(ocurrencias, orden)

        hoja.append(
            [
                caso["id"],
                caso["titulo"],
                caso["descripcion"],
                caso["tipoIa"],
                caso["tipoValor"],
                etiquetas.get(dominio, ""),
                capacidad,
                len(ocurrencias),
                len(dominios),
                ", ".join(etiquetas[d] for d in dominios),
                ", ".join(o[2] for o in ocurrencias),
                "Sí" if caso.get("poc") else "No",
                caso.get("origen", ""),
                "",
            ]
        )

    estilar_cabecera(hoja)

    for indice, (_, ancho) in enumerate(CABECERAS, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho

    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical="top")
            celda.font = Font(size=10)

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(CABECERAS))}1"

    return hoja


def escribir_listas(libro, catalogo):
    """La hoja oculta que alimenta los desplegables de las dos etiquetas."""
    hoja = libro.create_sheet(HOJA_LISTAS)

    tipos_de_ia = [entrada["valor"] for entrada in catalogo["tiposDeIa"]]
    tipos_de_valor = [entrada["valor"] for entrada in catalogo["tiposDeValor"]]
    si_no = ["Sí", "No"]

    hoja.append(["Tipo de IA", "Tipo de valor", "Sí/No"])

    for indice in range(max(len(tipos_de_ia), len(tipos_de_valor), len(si_no))):
        hoja.append(
            [
                tipos_de_ia[indice] if indice < len(tipos_de_ia) else None,
                tipos_de_valor[indice] if indice < len(tipos_de_valor) else None,
                si_no[indice] if indice < len(si_no) else None,
            ]
        )

    hoja.sheet_state = "hidden"

    return len(tipos_de_ia), len(tipos_de_valor)


def anadir_desplegables(hoja_casos, filas_ia, filas_valor, total_casos):
    """Los desplegables, para que una correccion a mano use un valor de la lista.

    Sin ellos, corregir una etiqueta en el Excel y escribirla distinta —una
    tilde, un plural— rompe el cruce al traerla de vuelta al JSON.
    """
    ultima = total_casos + 1

    validaciones = [
        (f"=Listas!$A$2:$A${filas_ia + 1}", f"D2:D{ultima}"),
        (f"=Listas!$B$2:$B${filas_valor + 1}", f"E2:E{ultima}"),
        ("=Listas!$C$2:$C$3", f"L2:L{ultima}"),
    ]

    for formula, rango in validaciones:
        validacion = DataValidation(type="list", formula1=formula, allow_blank=True)
        hoja_casos.add_data_validation(validacion)
        validacion.add(rango)


def escribir_resumen(libro, catalogo, etiquetas, apariciones):
    hoja = libro.create_sheet("3. Resumen por dominio")

    tipos_de_valor = [entrada["valor"] for entrada in catalogo["tiposDeValor"]]
    tipos_de_ia = [entrada["valor"] for entrada in catalogo["tiposDeIa"]]

    hoja.append(["Dominio"] + tipos_de_valor + ["Total", "Con POC"])

    orden = list(etiquetas)

    conteo = {did: Counter() for did in orden}
    con_poc = Counter()
    totales = Counter()

    for caso in catalogo["casos"]:
        dominio, _ = principal(apariciones.get(caso["titulo"], []), orden)

        if not dominio:
            continue

        conteo[dominio][caso["tipoValor"]] += 1
        totales[dominio] += 1

        if caso.get("poc"):
            con_poc[dominio] += 1

    for did in orden:
        if not totales[did]:
            continue

        hoja.append(
            [etiquetas[did]]
            + [conteo[did][tipo] for tipo in tipos_de_valor]
            + [totales[did], con_poc[did]]
        )

    hoja.append(
        ["TOTAL"]
        + [sum(conteo[d][tipo] for d in orden) for tipo in tipos_de_valor]
        + [sum(totales.values()), sum(con_poc.values())]
    )

    hoja.append([])
    hoja.append(["Por tipo de IA"])

    por_tipo_de_ia = Counter(caso["tipoIa"] for caso in catalogo["casos"])

    for tipo in tipos_de_ia:
        hoja.append([tipo, por_tipo_de_ia[tipo]])

    estilar_cabecera(hoja)

    hoja[f"A{len(totales) + 2}"].font = Font(bold=True)
    hoja[f"A{len(totales) + 4}"].font = Font(bold=True, color="3E6F11")

    hoja.column_dimensions["A"].width = 26

    for indice in range(2, len(tipos_de_valor) + 2):
        hoja.column_dimensions[get_column_letter(indice)].width = 21

    hoja.column_dimensions[get_column_letter(len(tipos_de_valor) + 2)].width = 9
    hoja.column_dimensions[get_column_letter(len(tipos_de_valor) + 3)].width = 11


def escribir_cambios(libro):
    hoja = libro.create_sheet("4. Cambios aplicados")

    for fila in CAMBIOS_APLICADOS:
        hoja.append(list(fila))

    estilar_cabecera(hoja)

    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical="top")
            celda.font = Font(size=10)

    hoja.column_dimensions["A"].width = 14
    hoja.column_dimensions["B"].width = 58
    hoja.column_dimensions["C"].width = 96


def escribir_leyenda(libro, catalogo):
    hoja = libro.create_sheet("5. Leyenda")

    hoja.append(["Tipo de IA"])

    for entrada in catalogo["tiposDeIa"]:
        hoja.append([entrada["valor"], entrada["definicion"]])

    hoja.append([])
    hoja.append(["Tipo de valor"])

    for entrada in catalogo["tiposDeValor"]:
        hoja.append([entrada["valor"], entrada["definicion"]])

    for etiqueta, texto in CRITERIOS_Y_ADVERTENCIA:
        hoja.append([etiqueta, texto])

    for fila in hoja.iter_rows():
        etiqueta, texto = fila[0], (fila[1] if len(fila) > 1 else None)

        if etiqueta.value and not (texto and texto.value):
            etiqueta.font = Font(bold=True, size=11, color=AZUL_DELOITTE)
        elif etiqueta.value:
            etiqueta.font = Font(bold=True, size=10)

        if texto:
            texto.alignment = Alignment(wrap_text=True, vertical="top")

    hoja.column_dimensions["A"].width = 30
    hoja.column_dimensions["B"].width = 116


def main():
    salida = Path(sys.argv[1]) if len(sys.argv) > 1 else SALIDA_POR_DEFECTO

    catalogo = leer_catalogo()
    etiquetas, apariciones = leer_apariciones()

    sin_apariciones = [
        caso["titulo"] for caso in catalogo["casos"] if not apariciones.get(caso["titulo"])
    ]

    if sin_apariciones:
        # No es fatal —el documento se puede generar igual— pero si se calla, la
        # hoja de resumen sale descuadrada y nadie sabe por que.
        print(
            f"AVISO: {len(sin_apariciones)} caso(s) del catalogo no aparecen en ninguna "
            "subcapacidad; saldran sin dominio. Ejecuta scripts/check_domains_sync.py."
        )

    libro = Workbook()

    escribir_como_usar(libro)

    hoja_casos = escribir_casos(libro, catalogo, etiquetas, apariciones)

    escribir_resumen(libro, catalogo, etiquetas, apariciones)
    escribir_cambios(libro)
    escribir_leyenda(libro, catalogo)

    filas_ia, filas_valor = escribir_listas(libro, catalogo)

    anadir_desplegables(hoja_casos, filas_ia, filas_valor, len(catalogo["casos"]))

    libro.save(salida)

    apariciones_totales = sum(len(v) for v in apariciones.values())

    print(
        f"OK: {salida} -> {len(catalogo['casos'])} casos, "
        f"{apariciones_totales} apariciones en {len(etiquetas)} dominios"
    )


if __name__ == "__main__":
    main()
